"""Generate the architecture's standard input workbook, independently of analysis."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .reader import REQUIRED_COLUMNS

TEMPLATE_COLUMNS = {
    "Materials": (
        *REQUIRED_COLUMNS["Materials"],
        "supplier",
        "grade",
        "notes",
        "density",
        "density_kind",
    ),
    "Measurements": (*REQUIRED_COLUMNS["Measurements"], "version"),
    "PSD": REQUIRED_COLUMNS["PSD"],
    "Recipe": (*REQUIRED_COLUMNS["Recipe"], "recipe_name"),
}


def generate_excel_template(path: str | Path, *, include_example: bool = True) -> Path:
    """Write .xlsx; optional clearly marked fictional rows form a valid round-trip example."""
    destination = Path(path)
    workbook = Workbook()
    instructions = workbook.active
    assert instructions is not None
    instructions.title = "说明"
    for row in (
        ("PSD Mix Analyzer — INPUT TEMPLATE", "Version 1"),
        ("EXAMPLE DATA / 示例数据", "Fictional data only. Replace all example rows before use."),
        ("Particle size", "μm; positive and strictly increasing within each measurement"),
        ("Passing and mass fractions", "0–100 numbers (25), text 25%, or native Excel 25%"),
        ("Density", "kg/m³; optional, supply density_kind together: particle / true / bulk"),
        ("Sheets", "Materials → Measurements → PSD; Recipe refers to material_code"),
        ("Blank optional values", "supplier / grade / notes / density may be blank"),
        ("Input rules", "No formulas, no automatic sorting, deletion, normalization or smoothing"),
        ("Measurements", "basis is explicit; method and protocol preserve measurement provenance"),
        (
            "Recipe",
            "Imported versions are drafts; importing never selects a batch or changes production",
        ),
        (
            "Identity fields",
            "Use text identifiers; preserve batch_key and measurement_key uniqueness",
        ),
    ):
        instructions.append(row)
    for name, columns in TEMPLATE_COLUMNS.items():
        sheet = workbook.create_sheet(name)
        sheet.append(columns)
    if include_example:
        for suffix in ("A", "B"):
            workbook["Materials"].append(
                (
                    f"EX-{suffix}",
                    f"Example material {suffix}",
                    f"EX-BATCH-{suffix}",
                    f"example-{suffix}",
                    None,
                    None,
                    "EXAMPLE ONLY",
                    None,
                    None,
                )
            )
            workbook["Measurements"].append(
                (
                    f"EX-PSD-{suffix}",
                    f"EX-BATCH-{suffix}",
                    "mass",
                    "example method",
                    "example protocol",
                    "1",
                )
            )
            for size, passing in ((1, 0), (16, 25), (81, 50), (256, 75), (625, 100)):
                workbook["PSD"].append((f"EX-PSD-{suffix}", size, passing))
        for suffix, fraction in (("A", 25), ("B", 75)):
            workbook["Recipe"].append(
                ("EX-RECIPE", "v1", f"EX-LINE-{suffix}", f"EX-{suffix}", fraction, "EXAMPLE ONLY")
            )
    for sheet in workbook:
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="234E70")
        if sheet.title != "说明":
            sheet.auto_filter.ref = sheet.dimensions
        for number, cells in enumerate(sheet.columns, 1):
            longest = max((len(str(c.value)) for c in cells if c.value is not None), default=10)
            sheet.column_dimensions[get_column_letter(number)].width = min(max(longest + 3, 15), 95)
    try:
        workbook.save(destination)
    finally:
        workbook.close()
    return destination
