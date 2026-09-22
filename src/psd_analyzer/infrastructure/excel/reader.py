"""Workbook I/O and schema inspection, with no particle-distribution calculations."""

from dataclasses import dataclass
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .records import CellValue, ExcelImportError, ImportIssue, SheetRow

REQUIRED_COLUMNS: dict[str, tuple[str, ...]] = {
    "Materials": ("material_code", "material_name", "batch_key", "batch_no"),
    "Measurements": ("measurement_key", "batch_key", "basis", "method", "protocol"),
    "PSD": ("measurement_key", "particle_size_um", "cumulative_passing_pct"),
    "Recipe": ("recipe_code", "recipe_version", "line_key", "material_code", "mass_fraction_pct"),
}


@dataclass(frozen=True)
class WorkbookData:
    filename: str
    source_hash: str
    sheets: tuple[tuple[str, tuple[tuple[CellValue, ...], ...]], ...]


class WorkbookReader:
    """Read literals and formats, retaining formulas as errors rather than cached values."""

    def read(self, path: str | Path) -> WorkbookData:
        filename = str(path)
        try:
            raw = Path(path).read_bytes()
            workbook = load_workbook(BytesIO(raw), read_only=True, data_only=False)
            try:
                sheets = tuple(
                    (
                        name,
                        tuple(
                            tuple(
                                CellValue(c.value, c.number_format, c.data_type == "f") for c in row
                            )
                            for row in workbook[name].iter_rows()
                        ),
                    )
                    for name in workbook.sheetnames
                    if name in REQUIRED_COLUMNS
                )
            finally:
                workbook.close()
        except (
            OSError,
            ValueError,
            KeyError,
            BadZipFile,
            InvalidFileException,
            EOFError,
            SyntaxError,
        ) as exc:
            raise ExcelImportError(
                (ImportIssue(filename, "", None, None, None, "workbook_unreadable", str(exc)),)
            ) from exc
        return WorkbookData(filename, sha256(raw).hexdigest(), sheets)


class ExcelSchemaValidator:
    """Resolve exact column names once, before parsing any data rows."""

    def validate(self, workbook: WorkbookData) -> tuple[SheetRow, ...]:
        issues: list[ImportIssue] = []
        rows: list[SheetRow] = []
        available = dict(workbook.sheets)
        for sheet, required in REQUIRED_COLUMNS.items():
            if sheet not in available:
                issues.append(
                    ImportIssue(
                        workbook.filename,
                        sheet,
                        None,
                        None,
                        None,
                        "missing_sheet",
                        f"Required sheet {sheet} is missing",
                    )
                )
                continue
            cells = available[sheet]
            headers = tuple(
                str(c.value).strip() if c.value is not None else ""
                for c in (cells[0] if cells else ())
            )
            for column in required:
                if column not in headers:
                    issues.append(
                        ImportIssue(
                            workbook.filename,
                            sheet,
                            1,
                            column,
                            None,
                            "missing_column",
                            f"Required column {column} is missing",
                        )
                    )
            duplicates = {h for h in headers if h and headers.count(h) > 1}
            for column in sorted(duplicates):
                issues.append(
                    ImportIssue(
                        workbook.filename,
                        sheet,
                        1,
                        column,
                        column,
                        "duplicate_column",
                        "Column name is repeated",
                    )
                )
            if duplicates:
                continue
            for number, values in enumerate(cells[1:], 2):
                if all(
                    c.value is None or (isinstance(c.value, str) and not c.value.strip())
                    for c in values
                ):
                    continue
                rows.append(
                    SheetRow(
                        sheet,
                        number,
                        tuple((h, value) for h, value in zip(headers, values, strict=False) if h),
                    )
                )
        if issues:
            raise ExcelImportError(tuple(issues))
        return tuple(rows)
