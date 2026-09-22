"""A hand-calculable workbook exercises public input and numerical interfaces."""

import pytest
from openpyxl import Workbook

from psd_analyzer.domain.models.analysis_profile import AnalysisProfile
from psd_analyzer.domain.models.analysis_result import FitStatus
from psd_analyzer.domain.models.psd import PSD
from psd_analyzer.domain.models.recipe import RecipeStatus
from psd_analyzer.domain.services.interpolation import LogSizeLinearInterpolator
from psd_analyzer.domain.services.losses import SquaredErrorLoss
from psd_analyzer.domain.services.metrics import calculate_metrics
from psd_analyzer.domain.services.packing_models import ModifiedAndreasen
from psd_analyzer.domain.services.psd_mixing import PSDMixingService
from psd_analyzer.domain.services.q_fitting import fit_q
from psd_analyzer.domain.services.queries import evaluate_at_sizes
from psd_analyzer.infrastructure.excel import import_excel_workbook


def test_excel_to_mix_fit_and_metrics(tmp_path):
    # Fourth roots of D are 1,2,3,4,5; hence q=.25 gives P=(root-1)/4.
    sizes = (1, 16, 81, 256, 625)
    percentages = (0, 25, 50, 75, 100)
    book = Workbook()
    book.remove(book.active)
    materials = book.create_sheet("Materials")
    materials.append(["material_code", "material_name", "batch_key", "batch_no"])
    materials.append(["A", "Example A", "batch-A", "example-1"])
    materials.append(["B", "Example B", "batch-B", "example-2"])
    measurements = book.create_sheet("Measurements")
    measurements.append(["measurement_key", "batch_key", "basis", "method", "protocol"])
    measurements.append(["PSD-A", "batch-A", "mass", "example", "example-v1"])
    measurements.append(["PSD-B", "batch-B", "mass", "example", "example-v1"])
    psd_sheet = book.create_sheet("PSD")
    psd_sheet.append(["measurement_key", "particle_size_um", "cumulative_passing_pct"])
    for measurement_key in ("PSD-A", "PSD-B"):
        for size, passing in zip(sizes, percentages, strict=True):
            psd_sheet.append([measurement_key, size, passing])
    recipes = book.create_sheet("Recipe")
    recipes.append(
        ["recipe_code", "recipe_version", "line_key", "material_code", "mass_fraction_pct"]
    )
    recipes.append(["R", "1", "line-A", "A", 25])
    recipes.append(["R", "1", "line-B", "B", 75])
    path = tmp_path / "hand_calculable.xlsx"
    book.save(path)
    book.close()

    imported = import_excel_workbook(path)
    recipe = imported.recipes[0]
    assert recipe.status is RecipeStatus.DRAFT
    assert tuple(line.mass_fraction for line in recipe.lines) == (0.25, 0.75)
    assert {batch.supplier for batch in imported.batches} == {""}
    # Selection is explicit in this caller; importing never chooses a batch for a recipe.
    measurement_by_id = {measurement.psd_id: measurement for measurement in imported.measurements}
    selected = {"A": measurement_by_id["PSD-A"], "B": measurement_by_id["PSD-B"]}
    interpolator = LogSizeLinearInterpolator()
    mixed = PSDMixingService().mix(
        tuple(selected[line.material_id].psd for line in recipe.lines),
        tuple(line.mass_fraction for line in recipe.lines),
        interpolator,
    )
    assert isinstance(mixed, PSD)
    assert mixed.particle_size_um == sizes
    assert mixed.cumulative_passing == pytest.approx((0, 0.25, 0.5, 0.75, 1))

    profile = AnalysisProfile("example", "1", 1, 625, key_sizes_um=(16, 81))
    model = ModifiedAndreasen()
    fit = fit_q(
        mixed.particle_size_um,
        mixed.cumulative_passing,
        profile=profile,
        model=model,
        loss=SquaredErrorLoss(),
    )
    assert fit.status is FitStatus.SUCCESS
    assert fit.converged and fit.point_count == 5
    assert fit.equivalent_q == pytest.approx(0.25, abs=1e-6)
    target = model.cumulative_passing(sizes, q=0.25, d_min_um=1, d_max_um=625)
    metrics = calculate_metrics(mixed.cumulative_passing, target)
    assert metrics.sse < 1e-25
    assert fit.metrics.sse < 1e-12
    assert evaluate_at_sizes(mixed, (16, 81), interpolator).cumulative_passing == (0.25, 0.5)
