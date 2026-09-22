"""Real Excel bytes → application → SQLite snapshots → existing figure builders."""

from datetime import UTC, datetime, timedelta
from io import BytesIO
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from psd_analyzer.application.dto.history import AnalysisType, HistoryFilter
from psd_analyzer.ui.streamlit_app.composition import build_container
from psd_analyzer.visualization import build_psd_figure, build_q_trend_figure


def load_approved_example(container):
    """Fixture explicitly represents an external approval; product UI has no release action."""
    book = load_workbook(BytesIO(container.generate_template.execute()))
    sheet = book["Recipe"]
    columns = {cell.value: cell.column for cell in sheet[1]}
    for row in (2, 3):
        sheet.cell(row, columns["status"], "released")
        sheet.cell(row, columns["approval_reference"], "EXAMPLE-EXTERNAL-APPROVAL")
    output = BytesIO()
    book.save(output)
    book.close()
    return container.import_workbook.execute(output.getvalue(), "approved-example.xlsx")


def test_complete_excel_application_history_visualization_pipeline(tmp_path):
    path = tmp_path / "history.db"
    container = build_container(path)
    workbook = load_approved_example(container)
    recipe = workbook.recipes[0]
    choices = {"EX-LINE-A": "EX-PSD-A", "EX-LINE-B": "EX-PSD-B"}
    profile = container.workflow.profile(1, 625, target_q=0.4, key_sizes="16,81,256")
    result = container.workflow.production(workbook, recipe, choices, profile)
    assert result.equivalent_q == pytest.approx(0.25, abs=1e-6)
    assert result.mixed_curve.cumulative_passing == (0, 0.25, 0.5, 0.75, 1)
    assert container.list_history.execute() == ()  # Calculation never saves implicitly.
    when = datetime(2026, 9, 22, 10, tzinfo=UTC)
    identifier = str(uuid4())
    saved = container.save_analysis.execute(
        result, AnalysisType.PRODUCTION_MONITORING, analysis_id=identifier, created_at=when
    )
    # New composition demonstrates persistence across UI reruns/process lifetimes.
    fresh = build_container(path)
    restored = fresh.get_detail.execute(identifier)
    assert restored == saved
    assert restored.result == result
    assert restored.current.components == result.components
    assert restored.current.key_passing == result.key_passing
    figure = build_psd_figure(restored.result)
    assert len(figure.data) >= 2
    assert any(tuple(trace.y) == (0, 25, 50, 75, 100) for trace in figure.data)
    fresh.set_baseline.execute(identifier)
    assert fresh.get_baseline.execute(recipe.recipe_id, recipe.version).analysis_id == identifier
    current = fresh.workflow.production(workbook, recipe, choices, profile)
    comparison = fresh.workflow.compare_baseline(restored.current, current)
    assert comparison.delta_q == 0
    assert comparison.metrics.rmse == 0
    next_id = str(uuid4())
    fresh.save_analysis.execute(
        current,
        AnalysisType.PRODUCTION_MONITORING,
        analysis_id=next_id,
        created_at=when + timedelta(days=1),
    )
    historical = fresh.compare_history.execute(identifier, next_id)
    assert historical == comparison
    summaries = fresh.list_history.execute(HistoryFilter(recipe_id=recipe.recipe_id))
    assert len(summaries) == 2
    trend = build_q_trend_figure(summaries)
    assert len(trend.data) == 2  # q and optional target reference, no SPC.
    assert tuple(trend.data[0].x) == (when, when + timedelta(days=1))


def test_substitution_and_simulation_are_saved_as_distinct_types(tmp_path):
    container = build_container(tmp_path / "history.db")
    book = load_approved_example(container)
    recipe = book.recipes[0]
    choices = {"EX-LINE-A": "EX-PSD-A", "EX-LINE-B": "EX-PSD-B"}
    profile = container.workflow.profile(1, 625)
    comparison = container.workflow.substitute(
        book, recipe, choices, {"EX-LINE-A": "EX-PSD-B"}, profile
    )
    saved = container.save_analysis.execute(
        comparison, AnalysisType.MATERIAL_SUBSTITUTION, analysis_id=str(uuid4())
    )
    restored = container.get_detail.execute(saved.analysis_id)
    assert restored.result == comparison
    assert restored.result.baseline.components[0].batch.material_id == "EX-A"
    assert restored.result.current.components[0].batch.material_id == "EX-B"
    assert len(build_psd_figure(restored.result).data) == 2
    simulation = container.workflow.simulation(
        book, recipe, choices, {"EX-LINE-A": 50, "EX-LINE-B": 50}, profile
    )
    record = container.save_analysis.execute(
        simulation, AnalysisType.RECIPE_SIMULATION, analysis_id=str(uuid4())
    )
    assert container.get_detail.execute(record.analysis_id).current.is_simulation
    assert (
        len(
            container.list_history.execute(
                HistoryFilter(analysis_type=AnalysisType.RECIPE_SIMULATION)
            )
        )
        == 1
    )
