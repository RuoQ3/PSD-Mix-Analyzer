"""Input-only adapter tests use actual small workbooks, never numerical model code."""

from dataclasses import FrozenInstanceError
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import load_workbook

from psd_analyzer.infrastructure.excel import (
    ExcelImportError,
    generate_excel_template,
    import_excel_workbook,
)
from psd_analyzer.infrastructure.excel.parser import SheetParser
from psd_analyzer.infrastructure.excel.records import CellValue, SheetRow


@pytest.fixture
def workbook_path(tmp_path):
    return generate_excel_template(tmp_path / "input.xlsx")


def edit(path, sheet, column, row, value, number_format=None):
    workbook = load_workbook(path)
    worksheet = workbook[sheet]
    columns = {cell.value: cell.column for cell in worksheet[1]}
    cell = worksheet.cell(row, columns[column], value)
    if value is None:
        cell.value = None
    if number_format is not None:
        cell.number_format = number_format
    workbook.save(path)
    workbook.close()


def failures(path):
    with pytest.raises(ExcelImportError) as raised:
        import_excel_workbook(path)
    assert all(issue.file == str(path) for issue in raised.value.issues)
    return raised.value.issues


def test_success_and_template_round_trip(workbook_path):
    result = import_excel_workbook(workbook_path)
    assert [m.material_id for m in result.materials] == ["EX-A", "EX-B"]
    assert [b.batch_id for b in result.batches] == ["EX-BATCH-A", "EX-BATCH-B"]
    assert result.measurements[0].psd_id == "EX-PSD-A"
    assert result.measurements[0].psd.particle_size_um == (1, 16, 81, 256, 625)
    assert result.measurements[0].psd.cumulative_passing == (0, 0.25, 0.5, 0.75, 1)
    assert result.recipes[0].lines[0].mass_fraction == 0.25
    assert result.measurements[0].source_hash == result.source_hash
    assert len(result.source_hash) == 64
    assert result.batches[0].supplier == ""
    assert result.recipes[0].status.value == "draft"
    with pytest.raises(FrozenInstanceError):
        result.source_hash = "changed"
    workbook = load_workbook(workbook_path)
    assert (
        "EXAMPLE" in workbook["说明"]["B2"].value.upper()
        or "Fictional" in workbook["说明"]["B2"].value
    )
    workbook.close()


@pytest.mark.parametrize(
    "raw,number_format", [(25, "General"), ("25%", "General"), (0.25, "0%"), (0.25, "0.00%")]
)
def test_percentage_encodings(workbook_path, raw, number_format):
    edit(workbook_path, "PSD", "cumulative_passing_pct", 3, raw, number_format)
    edit(workbook_path, "Recipe", "mass_fraction_pct", 2, raw, number_format)
    imported = import_excel_workbook(workbook_path)
    assert imported.measurements[0].psd.cumulative_passing[1] == 0.25
    assert imported.recipes[0].lines[0].mass_fraction == 0.25


def test_percentage_units_are_not_guessed(workbook_path):
    edit(workbook_path, "PSD", "cumulative_passing_pct", 3, 0.25)
    assert import_excel_workbook(workbook_path).measurements[0].psd.cumulative_passing[1] == 0.0025


def test_missing_sheet(workbook_path):
    workbook = load_workbook(workbook_path)
    del workbook["PSD"]
    workbook.save(workbook_path)
    workbook.close()
    issue = failures(workbook_path)[0]
    assert (issue.sheet, issue.code) == ("PSD", "missing_sheet")


def test_missing_and_duplicate_columns(workbook_path):
    workbook = load_workbook(workbook_path)
    workbook["PSD"]["B1"] = "measurement_key"
    workbook.save(workbook_path)
    workbook.close()
    assert {i.code for i in failures(workbook_path)} >= {"missing_column", "duplicate_column"}


@pytest.mark.parametrize(
    "sheet,column,row,value,code",
    [
        ("Recipe", "material_code", 2, "unknown", "unknown_material"),
        ("PSD", "measurement_key", 2, "unknown", "unknown_measurement"),
        ("Measurements", "batch_key", 2, "unknown", "unknown_batch"),
        ("PSD", "particle_size_um", 3, 1, "duplicate_particle_size"),
        ("PSD", "particle_size_um", 3, 0.5, "nonascending_particle_size"),
        ("PSD", "particle_size_um", 3, 0, "nonpositive_value"),
        ("PSD", "particle_size_um", 3, "bad", "invalid_number"),
        ("PSD", "particle_size_um", 3, "NaN", "invalid_number"),
        ("PSD", "particle_size_um", 3, "inf", "invalid_number"),
        ("PSD", "cumulative_passing_pct", 3, 101, "percentage_out_of_range"),
        ("PSD", "cumulative_passing_pct", 4, 20, "nonmonotone_passing"),
        ("Recipe", "mass_fraction_pct", 2, 20, "invalid_recipe_total"),
        ("Recipe", "mass_fraction_pct", 2, -5, "percentage_out_of_range"),
        ("Materials", "material_code", 2, None, "required_value"),
        ("PSD", "particle_size_um", 3, "=4*4", "formula_not_supported"),
        ("Measurements", "basis", 2, "mystery", "invalid_basis"),
        ("Materials", "density", 2, 1000, "density_pair"),
        ("Materials", "density_kind", 2, "mystery", "invalid_density_kind"),
    ],
)
def test_invalid_cells_have_precise_issues(workbook_path, sheet, column, row, value, code):
    edit(workbook_path, sheet, column, row, value)
    issue = next(i for i in failures(workbook_path) if i.code == code)
    assert issue.sheet == sheet
    assert issue.row == row
    if code not in ("density_pair",):
        assert issue.column == column
        assert issue.value == value
    assert issue.severity == "error"
    assert issue.message


def test_multiple_errors_are_reported_in_one_pass(workbook_path):
    edit(workbook_path, "PSD", "particle_size_um", 3, -1)
    edit(workbook_path, "PSD", "cumulative_passing_pct", 4, 200)
    edit(workbook_path, "Recipe", "material_code", 2, "unknown")
    issues = failures(workbook_path)
    assert {i.code for i in issues} >= {
        "nonpositive_value",
        "percentage_out_of_range",
        "unknown_material",
    }


def test_insufficient_psd(workbook_path):
    workbook = load_workbook(workbook_path)
    workbook["PSD"].delete_rows(3, 4)
    workbook.save(workbook_path)
    workbook.close()
    assert "insufficient_psd_points" in {i.code for i in failures(workbook_path)}


def test_optional_blanks_and_trimming(workbook_path):
    edit(workbook_path, "Materials", "supplier", 2, "   ")
    edit(workbook_path, "Materials", "notes", 2, None)
    edit(workbook_path, "Materials", "material_name", 2, "  Trimmed  ")
    result = import_excel_workbook(workbook_path)
    assert result.materials[0].name == "Trimmed"
    assert result.materials[0].notes == ""
    assert result.batches[0].supplier == ""


def test_duplicate_identities_and_conflicting_names(workbook_path):
    edit(workbook_path, "Materials", "batch_key", 3, "EX-BATCH-A")
    edit(workbook_path, "Measurements", "measurement_key", 3, "EX-PSD-A")
    edit(workbook_path, "Recipe", "line_key", 3, "EX-LINE-A")
    edit(workbook_path, "Recipe", "recipe_name", 3, "Different")
    assert {i.code for i in failures(workbook_path)} >= {
        "duplicate_batch",
        "duplicate_measurement",
        "duplicate_recipe_line",
        "conflicting_recipe_name",
    }


def test_material_can_have_several_batches(workbook_path):
    edit(workbook_path, "Materials", "material_code", 3, "EX-A")
    edit(workbook_path, "Materials", "material_name", 3, "Example material A")
    edit(workbook_path, "Recipe", "material_code", 3, "EX-A")
    imported = import_excel_workbook(workbook_path)
    assert len(imported.materials) == 1
    assert len(imported.batches) == 2
    assert len(imported.measurements) == 2


def test_material_metadata_cannot_conflict(workbook_path):
    edit(workbook_path, "Materials", "material_code", 3, "EX-A")
    assert "conflicting_material" in {i.code for i in failures(workbook_path)}


def test_no_rows_and_blank_rows(tmp_path, workbook_path):
    blank = generate_excel_template(tmp_path / "empty.xlsx", include_example=False)
    assert "empty_sheet" in {i.code for i in failures(blank)}
    workbook = load_workbook(workbook_path)
    workbook["PSD"].append((None, None, " "))
    workbook.save(workbook_path)
    workbook.close()
    assert len(import_excel_workbook(workbook_path).measurements) == 2


def test_unreadable_file_and_malformed_xml(tmp_path, workbook_path):
    invalid = tmp_path / "invalid.xlsx"
    invalid.write_text("not an Excel workbook")
    assert failures(invalid)[0].code == "workbook_unreadable"
    assert failures(tmp_path / "missing.xlsx")[0].code == "workbook_unreadable"
    with ZipFile(workbook_path) as original, ZipFile(invalid, "w", ZIP_DEFLATED) as output:
        for item in original.infolist():
            output.writestr(
                item,
                b"<malformed"
                if item.filename == "xl/workbook.xml"
                else original.read(item.filename),
            )
    assert failures(invalid)[0].code == "workbook_unreadable"


@pytest.mark.parametrize("value", [float("nan"), float("inf"), True])
def test_parser_rejects_nonfinite_and_boolean_values(value):
    issues = []
    parser = SheetParser("in-memory.xlsx", issues)
    row = SheetRow("PSD", 9, (("particle_size_um", CellValue(value)),))
    assert parser.number(row, "particle_size_um") is None
    assert issues[0].row == 9
    assert issues[0].code == "invalid_number"


def test_infrastructure_does_not_import_numerical_services():
    import ast

    root = Path(__file__).parents[2] / "src/psd_analyzer/infrastructure/excel"
    for path in root.glob("*.py"):
        for node in ast.walk(ast.parse(path.read_text())):
            if isinstance(node, ast.ImportFrom):
                module = node.module or ""
                assert "services" not in module
                assert "scipy" not in module
                assert "packing" not in module
