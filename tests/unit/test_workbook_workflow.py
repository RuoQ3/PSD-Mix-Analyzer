"""Presentation-facing orchestration keeps recipe invariants in the existing domain."""

from dataclasses import replace
from io import BytesIO

import pytest
from openpyxl import load_workbook

from psd_analyzer.application.dto.workbook import ExcelImportError, ImportedWorkbook
from psd_analyzer.application.exceptions import InvalidAnalysisConfiguration, MissingMaterialPSD
from psd_analyzer.application.use_cases.workbook_analysis import WorkbookAnalysis
from psd_analyzer.domain.exceptions import (
    DomainValidationError,
    IncompatibleComparisonError,
    RecipeLockedError,
)
from psd_analyzer.domain.models.recipe import RecipeStatus
from psd_analyzer.infrastructure.excel.gateway import ExcelWorkbookGateway
from psd_analyzer.infrastructure.excel.records import ImportedWorkbook as LegacyImportedWorkbook


@pytest.fixture
def imported():
    gateway = ExcelWorkbookGateway()
    return gateway.import_bytes(gateway.template_bytes(), "example.xlsx")


def selections(workbook):
    return {
        line.line_id: next(
            m.psd_id
            for m in workbook.measurements
            if next(b for b in workbook.batches if b.batch_id == m.batch_id).material_id
            == line.material_id
        )
        for line in workbook.recipes[0].lines
    }


def released(workbook):
    return replace(
        workbook.recipes[0],
        status=RecipeStatus.RELEASED,
        approval_reference="test-external-approval",
    )


def test_public_import_result_preserves_legacy_identity(imported):
    assert ImportedWorkbook is LegacyImportedWorkbook
    assert isinstance(imported, LegacyImportedWorkbook)


def test_workflow_calls_production_with_fixed_fractions(imported):
    service = WorkbookAnalysis()
    recipe = released(imported)
    result = service.production(imported, recipe, selections(imported), service.profile(1, 625))
    assert result.equivalent_q == pytest.approx(0.25, abs=1e-6)
    assert result.recipe == recipe
    assert result.actual_weights == (0.25, 0.75)
    assert result.target_curve is None
    assert result.fit_metrics.rmse < 1e-7
    assert not result.is_simulation


def test_draft_cannot_be_used_in_production(imported):
    service = WorkbookAnalysis()
    with pytest.raises(RecipeLockedError):
        service.production(
            imported, imported.recipes[0], selections(imported), service.profile(1, 625)
        )


def test_simulation_is_distinct_and_does_not_change_source(imported):
    service = WorkbookAnalysis()
    recipe = released(imported)
    result = service.simulation(
        imported,
        recipe,
        selections(imported),
        {line.line_id: 50 for line in recipe.lines},
        service.profile(1, 625),
    )
    assert result.is_simulation
    assert result.recipe.status == RecipeStatus.DRAFT
    assert result.recipe.recipe_id != recipe.recipe_id
    assert result.actual_weights == (0.5, 0.5)
    assert tuple(line.mass_fraction for line in recipe.lines) == (0.25, 0.75)
    assert recipe.status == RecipeStatus.RELEASED


@pytest.mark.parametrize(
    "values", [(50, 45), (-5, 105), (float("nan"), 100), (float("inf"), 0), (), (True, 99)]
)
def test_invalid_simulation_percentages_are_rejected(values):
    with pytest.raises(InvalidAnalysisConfiguration):
        WorkbookAnalysis.validate_simulation_fractions(values)


def test_valid_float_percentages():
    WorkbookAnalysis.validate_simulation_fractions((10, 20, 70))


def test_missing_or_unknown_selection(imported):
    service = WorkbookAnalysis()
    for choice in ({}, {"unknown": "EX-PSD-A"}, {"EX-LINE-A": "missing", "EX-LINE-B": "EX-PSD-B"}):
        with pytest.raises(MissingMaterialPSD):
            service.production(imported, released(imported), choice, service.profile(1, 625))


def test_substitution_uses_explicit_line_and_preserves_fractions(imported):
    service = WorkbookAnalysis()
    result = service.substitute(
        imported,
        released(imported),
        selections(imported),
        {"EX-LINE-A": "EX-PSD-B"},
        service.profile(1, 625),
    )
    assert result.baseline.components[0].batch.material_id == "EX-A"
    assert result.current.components[0].batch.material_id == "EX-B"
    assert result.current.actual_weights == result.baseline.actual_weights == (0.25, 0.75)
    assert result.metrics.rmse == 0
    assert result.delta_q == 0


@pytest.mark.parametrize(
    "kwargs",
    [
        {"key_sizes": "bad"},
        {"key_sizes": "10,10"},
        {"key_sizes": "45,10"},
        {"q_min": 0},
        {"interpolation": "bad"},
    ],
)
def test_invalid_profile_controls(kwargs):
    with pytest.raises((InvalidAnalysisConfiguration, DomainValidationError)):
        WorkbookAnalysis.profile(1, 625, **kwargs)


def test_baseline_comparison_has_no_silent_recalculation(imported):
    service = WorkbookAnalysis()
    result = service.production(
        imported, released(imported), selections(imported), service.profile(1, 625)
    )
    assert service.compare_baseline(result, result).metrics.rmse == 0
    changed = replace(result, profile=replace(result.profile, target_q=0.4))
    with pytest.raises(IncompatibleComparisonError):
        service.compare_baseline(result, changed)


def test_gateway_errors_keep_original_filename():
    with pytest.raises(ExcelImportError) as caught:
        ExcelWorkbookGateway().import_bytes(b"bad workbook", "user-file.xlsx")
    assert caught.value.issues
    assert {issue.file for issue in caught.value.issues} == {"user-file.xlsx"}


def recipe_workbook(status, approval, *, conflict=False):
    gateway = ExcelWorkbookGateway()
    workbook = load_workbook(BytesIO(gateway.template_bytes()))
    sheet = workbook["Recipe"]
    columns = {cell.value: cell.column for cell in sheet[1]}
    for row in (2, 3):
        sheet.cell(row, columns["status"], status)
        sheet.cell(row, columns["approval_reference"], approval)
    if conflict:
        sheet.cell(3, columns["approval_reference"], "different")
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return gateway.import_bytes(stream.getvalue(), "approved.xlsx")


def test_import_explicit_external_approval():
    workbook = recipe_workbook("released", "external-QA-123")
    assert workbook.recipes[0].status == RecipeStatus.RELEASED
    assert workbook.recipes[0].approval_reference == "external-QA-123"


@pytest.mark.parametrize(
    "status,approval,conflict,code",
    [
        ("released", "", False, "approval_required"),
        ("other", "", False, "invalid_recipe_status"),
        ("released", "QA-123", True, "conflicting_recipe_metadata"),
    ],
)
def test_import_invalid_release_metadata(status, approval, conflict, code):
    with pytest.raises(ExcelImportError) as caught:
        recipe_workbook(status, approval, conflict=conflict)
    assert code in {issue.code for issue in caught.value.issues}


def test_invalid_total_message_includes_actual_total():
    with pytest.raises(InvalidAnalysisConfiguration, match="95%"):
        WorkbookAnalysis.validate_simulation_fractions((25, 70))
